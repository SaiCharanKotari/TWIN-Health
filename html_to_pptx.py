import os
import time
import pathlib
from selenium import webdriver
from selenium.webdriver.edge.options import Options as EdgeOptions
from selenium.webdriver.chrome.options import Options as ChromeOptions
from pptx import Presentation
from pptx.util import Inches

def convert_html_to_pptx():
    workspace = pathlib.Path(__file__).parent.resolve()
    html_file = workspace / "FinalYearProjectpresentain.html"
    output_dir = workspace / "slide_images"
    output_pptx = workspace / "FinalYearProjectPresentation.pptx"
    
    output_dir.mkdir(exist_ok=True)
    
    html_uri = html_file.as_uri()
    print(f"Loading HTML: {html_uri}")
    
    opts = EdgeOptions()
    opts.add_argument("--headless=new")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--window-size=1920,1080")
    opts.add_argument("--hide-scrollbars")
    opts.add_argument("--force-device-scale-factor=2")  # 2x high-DPI crisp rendering
    
    try:
        driver = webdriver.Edge(options=opts)
    except Exception as e:
        print("Edge failed, trying Chrome:", e)
        copts = ChromeOptions()
        copts.add_argument("--headless=new")
        copts.add_argument("--disable-gpu")
        copts.add_argument("--window-size=1920,1080")
        copts.add_argument("--hide-scrollbars")
        copts.add_argument("--force-device-scale-factor=2")
        driver = webdriver.Chrome(options=copts)
    
    try:
        driver.set_window_size(1920, 1080)
        driver.get(html_uri)
        time.sleep(2)  # Wait for fonts & initial styles
        
        # Hide the floating interactive control buttons and print hint for clean PPT slides
        driver.execute_script("""
            const style = document.createElement('style');
            style.innerHTML = `
                .controls, .print-hint { display: none !important; }
                .slide { transition: none !important; }
            `;
            document.head.appendChild(style);
        """)
        
        # Get total slides count
        total_slides = driver.execute_script("return document.querySelectorAll('.slide').length;")
        print(f"Total slides found: {total_slides}")
        
        image_paths = []
        for i in range(total_slides):
            print(f"Capturing slide {i+1}/{total_slides}...")
            driver.execute_script(f"showSlide({i});")
            time.sleep(0.5)  # allow DOM update / rendering
            
            img_path = output_dir / f"slide_{i+1:02d}.png"
            driver.save_screenshot(str(img_path))
            image_paths.append(img_path)
            print(f"Saved: {img_path.name}")
            
    finally:
        driver.quit()
        
    print("\nCreating PowerPoint Presentation (.pptx)...")
    prs = Presentation()
    # 16:9 widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Blank layout is typically index 6
    blank_layout = None
    for layout in prs.slide_layouts:
        if len(layout.shapes) == 0:
            blank_layout = layout
            break
    if blank_layout is None:
        blank_layout = prs.slide_layouts[6]
        
    for i, img_path in enumerate(image_paths):
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            str(img_path),
            Inches(0),
            Inches(0),
            width=Inches(13.333),
            height=Inches(7.5)
        )
        print(f"Added slide {i+1} to PPTX")
        
    prs.save(str(output_pptx))
    print(f"\nSuccessfully generated: {output_pptx} with {len(image_paths)} slides!")

if __name__ == "__main__":
    convert_html_to_pptx()
